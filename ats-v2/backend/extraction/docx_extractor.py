# -*- coding: utf-8 -*-
"""
DOCX Extraction Module.

Extracts text from DOCX files while preserving paragraph structure.
"""

from __future__ import annotations

import io
import logging

import docx

log = logging.getLogger("ats.extraction.docx")

def extract_text(file_content: bytes) -> str:
    """
    Extracts text from a DOCX byte stream.
    
    Args:
        file_content: Raw bytes of the DOCX file.
        
    Returns:
        The extracted text as a single string.
    """
    log.info("Starting DOCX extraction")
    
    text_content = []
    
    try:
        doc = docx.Document(io.BytesIO(file_content))
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_content.append(text)
                
        # Also extract text from tables
        for table in doc.tables:
            text_content.append("--- Table ---")
            for row in table.rows:
                row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(row_data):
                    text_content.append(" | ".join(row_data))
            text_content.append("-------------")
            
    except Exception as e:
        log.error("DOCX extraction failed: %s", str(e))
        raise RuntimeError(f"DOCX extraction failed: {str(e)}") from e
        
    final_text = "\n".join(text_content).strip()
    log.debug("DOCX extracted %d characters", len(final_text))
    return final_text
